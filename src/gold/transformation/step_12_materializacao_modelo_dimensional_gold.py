import hashlib
import json
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


PROJECT_ROOT = Path(__file__).resolve().parents[3]
SILVER_PATH = (
    PROJECT_ROOT / "data" / "silver" / "food_commerce"
    / "silver_transactions_v3.parquet"
)
SPEC_PATH = (
    PROJECT_ROOT / "outputs" / "modeling" / "gold_v3"
    / "step_11_especificacao_modelo_dimensional.xlsx"
)
GOLD_DIR = PROJECT_ROOT / "data" / "gold" / "food_commerce" / "gold_v3"
REPORT_DIR = PROJECT_ROOT / "outputs" / "quality" / "gold_v3"
REPORT_PATH = REPORT_DIR / "step_12_materializacao_modelo_dimensional_gold.xlsx"
MANIFEST_DIR = PROJECT_ROOT / "outputs" / "manifests" / "gold_v3"
MANIFEST_PATH = MANIFEST_DIR / "step_12_materializacao_modelo_dimensional_gold.json"

EXPECTED_ROWS = 6495
EXPECTED_ORDERS = 2789
EXPECTED_CLIENTS = 330
EXPECTED_PRODUCTS = 36
TOLERANCE = 0.01

CLIENT_ATTRIBUTES = [
    "Idade", "Gênero", "Estado_Civil", "Escolaridade",
    "Situacao_Profissional", "Profissão", "Tempo_Experiencia_Anos",
    "Renda_Mensal_Cliente", "Ciclo_Vida_Cliente",
    "Intervalo_Medio_Compras_Dias", "Categoria_Favorita",
    "Forma_Pagamento_Preferida", "Dia_Semana_Preferido",
    "Ticket_Medio_Cliente", "Cidade", "Estado",
]

ORDER_ATTRIBUTES = [
    "Cliente_ID", "Data_Compra", "Horario_Compra", "Forma_Pagamento",
    "Campanha", "Canal_Marketing", "Evento_Externo", "Cidade", "Estado",
    "Frete", "Frete_Gratis", "Prazo_Entrega_Prometido",
    "Prazo_Entrega_Real", "Atraso_Entrega", "Avaliacao_Cliente",
    "Transportadora", "Pontos_Fidelidade", "Cupom_Utilizado",
]

UF_REGIAO = {
    "AC": "Norte", "AL": "Nordeste", "AP": "Norte", "AM": "Norte",
    "BA": "Nordeste", "CE": "Nordeste", "DF": "Centro-Oeste",
    "ES": "Sudeste", "GO": "Centro-Oeste", "MA": "Nordeste",
    "MT": "Centro-Oeste", "MS": "Centro-Oeste", "MG": "Sudeste",
    "PA": "Norte", "PB": "Nordeste", "PR": "Sul", "PE": "Nordeste",
    "PI": "Nordeste", "RJ": "Sudeste", "RN": "Nordeste", "RS": "Sul",
    "RO": "Norte", "RR": "Norte", "SC": "Sul", "SP": "Sudeste",
    "SE": "Nordeste", "TO": "Norte",
}


def utc_now():
    return datetime.now(timezone.utc)


def sha256_file(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def scalar(value):
    return None if pd.isna(value) else value


def stable_hash(values):
    text = "|".join("<NULL>" if pd.isna(v) else str(v) for v in values)
    return hashlib.sha256(text.encode("utf-8")).hexdigest().upper()


def normalize_text(series):
    return series.astype("string").str.strip()


def first_value(series):
    non_null = series.dropna()
    return non_null.iloc[0] if not non_null.empty else pd.NA


def validate_inputs(silver, specification):
    if len(silver) != EXPECTED_ROWS:
        raise ValueError(f"Silver deveria ter {EXPECTED_ROWS} linhas; tem {len(silver)}.")
    controls = {
        "pedidos": (silver["Pedido_ID"].nunique(), EXPECTED_ORDERS),
        "clientes": (silver["Cliente_ID"].nunique(), EXPECTED_CLIENTS),
        "produtos": (silver["Id_Item"].nunique(), EXPECTED_PRODUCTS),
    }
    failures = [f"{k}: {v[0]} != {v[1]}" for k, v in controls.items() if v[0] != v[1]]
    if failures:
        raise ValueError("Controles Silver divergentes: " + "; ".join(failures))

    generated_origins = {
        "gerada", "derivada", "calculada", "constante", "regra_negocio",
        "calendario", "calendario_oficial_2025", "dim_cliente", "dim_data",
        "fato_item_pedido", "fato_pedido", "derivada_UF", "Data_Hora_Compra",
    }
    required = set(
        specification.loc[
            ~specification["origem_silver"].isin(generated_origins),
            "origem_silver",
        ].dropna()
    )
    compound = {"Cidade+Estado"}
    required -= compound
    if "_bronze_source_row_number" in required:
        required.remove("_bronze_source_row_number")
        required.add("Bronze_Source_Row_Number")
    missing = sorted(required - set(silver.columns))
    if missing:
        raise ValueError(f"Colunas Silver exigidas pela Etapa 11 ausentes: {missing}")


def conflict_report(frame, key, attributes, object_name):
    rows = []
    for column in attributes:
        if column not in frame.columns:
            continue
        # Ausencia de valor nao constitui uma segunda categoria conflitante.
        # A REC016 reprova somente dois ou mais valores conhecidos diferentes.
        counts = frame.groupby(key, dropna=False)[column].nunique(dropna=True)
        for natural_key, quantity in counts[counts > 1].items():
            rows.append({
                "objeto": object_name,
                "chave": natural_key,
                "atributo": column,
                "valores_distintos": int(quantity),
            })
    return pd.DataFrame(rows, columns=["objeto", "chave", "atributo", "valores_distintos"])


def deterministic_dimension(series, sk, natural):
    values = normalize_text(series).drop_duplicates().sort_values(kind="stable")
    result = pd.DataFrame({natural: values.reset_index(drop=True)})
    result.insert(0, sk, np.arange(1, len(result) + 1, dtype=np.int64))
    return result


def build_date_dimension(calendar):
    dates = pd.date_range("2025-01-01", "2025-12-31", freq="D")
    dim = pd.DataFrame({"Data_Completa": dates})
    iso = dim["Data_Completa"].dt.isocalendar()
    dim.insert(0, "Data_SK", dim["Data_Completa"].dt.strftime("%Y%m%d").astype("int32"))
    dim["Ano"] = dim["Data_Completa"].dt.year.astype("int16")
    dim["Semestre"] = ((dim["Data_Completa"].dt.month - 1) // 6 + 1).astype("int8")
    dim["Trimestre"] = dim["Data_Completa"].dt.quarter.astype("int8")
    dim["Numero_Mes"] = dim["Data_Completa"].dt.month.astype("int8")
    months = {1:"Janeiro",2:"Fevereiro",3:"Março",4:"Abril",5:"Maio",6:"Junho",7:"Julho",8:"Agosto",9:"Setembro",10:"Outubro",11:"Novembro",12:"Dezembro"}
    weekdays = {0:"Segunda-feira",1:"Terça-feira",2:"Quarta-feira",3:"Quinta-feira",4:"Sexta-feira",5:"Sábado",6:"Domingo"}
    dim["Nome_Mes"] = dim["Numero_Mes"].map(months).astype("string")
    dim["Ano_Mes"] = dim["Data_Completa"].dt.strftime("%Y%m").astype("int32")
    dim["Dia_Mes"] = dim["Data_Completa"].dt.day.astype("int8")
    dim["Dia_Ano"] = dim["Data_Completa"].dt.dayofyear.astype("int16")
    dim["Semana_Ano"] = iso.week.astype("int8")
    dim["Numero_Dia_Semana"] = (dim["Data_Completa"].dt.weekday + 1).astype("int8")
    dim["Nome_Dia_Semana"] = dim["Data_Completa"].dt.weekday.map(weekdays).astype("string")
    dim["Final_Semana"] = dim["Data_Completa"].dt.weekday.ge(5)

    cal = calendar.rename(columns={
        "data": "Data_Completa", "nome_dia_especial": "Nome_Dia_Especial",
        "tipo_dia_especial": "Tipo_Dia_Especial", "periodo": "Periodo_Dia_Especial",
        "escopo": "Escopo_Dia_Especial", "fonte_normativa": "Fonte_Calendario",
    }).copy()
    cal["Data_Completa"] = pd.to_datetime(cal["Data_Completa"]).dt.normalize()
    dim = dim.merge(cal, on="Data_Completa", how="left", validate="1:1")
    dim["Feriado_Nacional"] = dim["Tipo_Dia_Especial"].eq("feriado_nacional")
    dim["Ponto_Facultativo_Federal"] = dim["Tipo_Dia_Especial"].eq("ponto_facultativo_federal")
    return dim


def build_time_dimension():
    minute = np.arange(1440, dtype=np.int16)
    hour = minute // 60
    result = pd.DataFrame({
        "Horario_SK": minute,
        "Horario": [f"{h:02d}:{m:02d}" for h, m in zip(hour, minute % 60)],
        "Hora": hour.astype(np.int8),
        "Minuto": (minute % 60).astype(np.int8),
    })
    result["Faixa_Horaria"] = pd.cut(
        result["Hora"], bins=[-1, 5, 11, 17, 23],
        labels=["Madrugada", "Manhã", "Tarde", "Noite"],
    ).astype("string")
    return result


def build_geography_dimension(silver):
    result = silver[["Cidade", "Estado"]].drop_duplicates().copy()
    result["Cidade"] = normalize_text(result["Cidade"])
    result["Estado"] = normalize_text(result["Estado"]).str.upper()
    result = result.sort_values(["Estado", "Cidade"], kind="stable").reset_index(drop=True)
    result.insert(0, "Geografia_SK", np.arange(1, len(result) + 1, dtype=np.int64))
    result["Regiao"] = result["Estado"].map(UF_REGIAO).astype("string")
    if result["Regiao"].isna().any():
        raise ValueError("Existem UFs sem mapeamento de região na dim_geografia.")
    result["Pais"] = "Brasil"
    return result


def build_client_dimension(silver, geography, load_time):
    source = silver.sort_values(["Cliente_ID", "Data_Compra", "Pedido_ID"], kind="stable")
    base = source.groupby("Cliente_ID", as_index=False).agg({c: first_value for c in CLIENT_ATTRIBUTES})
    base["Cidade"] = normalize_text(base["Cidade"])
    base["Estado"] = normalize_text(base["Estado"]).str.upper()
    base = base.merge(geography[["Geografia_SK", "Cidade", "Estado"]], on=["Cidade", "Estado"], how="left", validate="m:1")
    rename = {
        "Gênero": "Genero", "Profissão": "Profissao",
        "Ciclo_Vida_Cliente": "Ciclo_Vida_Informado",
        "Intervalo_Medio_Compras_Dias": "Intervalo_Medio_Informado",
        "Categoria_Favorita": "Categoria_Favorita_Informada",
        "Forma_Pagamento_Preferida": "Pagamento_Preferido_Informado",
        "Dia_Semana_Preferido": "Dia_Semana_Preferido_Informado",
        "Ticket_Medio_Cliente": "Ticket_Medio_Informado",
    }
    base = base.rename(columns=rename).drop(columns=["Cidade", "Estado"])
    base = base.sort_values("Cliente_ID", kind="stable").reset_index(drop=True)
    base.insert(0, "Cliente_SK", np.arange(1, len(base) + 1, dtype=np.int64))
    audit_exclusions = {"Cliente_SK", "Data_Carga", "Data_Atualizacao", "Hash_Atributos"}
    hash_columns = [c for c in base.columns if c not in audit_exclusions]
    base["Data_Carga"] = load_time
    base["Data_Atualizacao"] = load_time
    base["Hash_Atributos"] = base[hash_columns].apply(lambda row: stable_hash(row.tolist()), axis=1)
    return base


def build_product_dimension(silver):
    source = silver.sort_values(["Id_Item", "Pedido_ID"], kind="stable")
    result = source.groupby("Id_Item", as_index=False).agg({
        "Produto": first_value, "Categoria_Item": first_value, "Medida": first_value,
    }).sort_values("Id_Item", kind="stable").reset_index(drop=True)
    result.insert(0, "Produto_SK", np.arange(1, len(result) + 1, dtype=np.int64))
    result["Hash_Atributos"] = result[["Id_Item", "Produto", "Categoria_Item", "Medida"]].apply(lambda row: stable_hash(row.tolist()), axis=1)
    result["Registro_Ativo"] = True
    return result


def add_order_datetime(frame):
    date = pd.to_datetime(frame["Data_Compra"], errors="raise").dt.normalize()
    time_text = frame["Horario_Compra"].astype("string").str.extract(r"(\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?)", expand=False)
    delta = pd.to_timedelta(time_text, errors="raise")
    return date + delta


def map_foreign_keys(silver, dimensions):
    result = silver.copy()
    result["Data_Compra"] = pd.to_datetime(result["Data_Compra"], errors="raise").dt.normalize()
    result["Data_SK"] = result["Data_Compra"].dt.strftime("%Y%m%d").astype("int32")
    result["Data_Hora_Compra"] = add_order_datetime(result)
    result["Horario_SK"] = (result["Data_Hora_Compra"].dt.hour * 60 + result["Data_Hora_Compra"].dt.minute).astype("int16")
    result["Cidade"] = normalize_text(result["Cidade"])
    result["Estado"] = normalize_text(result["Estado"]).str.upper()
    mappings = [
        ("Cliente_ID", dimensions["dim_cliente"], "Cliente_ID", "Cliente_SK"),
        ("Id_Item", dimensions["dim_produto"], "Id_Item", "Produto_SK"),
        ("Forma_Pagamento", dimensions["dim_pagamento"], "Forma_Pagamento", "Pagamento_SK"),
        ("Campanha", dimensions["dim_campanha"], "Campanha", "Campanha_SK"),
        ("Canal_Marketing", dimensions["dim_canal_marketing"], "Canal_Marketing", "Canal_Marketing_SK"),
        ("Evento_Externo", dimensions["dim_evento_externo"], "Evento_Externo", "Evento_Externo_SK"),
        ("Transportadora", dimensions["dim_transportadora"], "Transportadora", "Transportadora_SK"),
    ]
    for source_column, dimension, natural, surrogate in mappings:
        result[source_column] = normalize_text(result[source_column])
        result = result.merge(dimension[[natural, surrogate]], left_on=source_column, right_on=natural, how="left", validate="m:1")
        if natural != source_column:
            result = result.drop(columns=[natural])
    result = result.merge(
        dimensions["dim_geografia"][["Cidade", "Estado", "Geografia_SK"]],
        on=["Cidade", "Estado"], how="left", validate="m:1",
    )
    return result


def build_item_fact(mapped):
    result = mapped.copy()
    result = result.sort_values(["Pedido_ID", "Bronze_Source_Row_Number"], kind="stable").reset_index(drop=True)
    result.insert(0, "Item_Pedido_SK", np.arange(1, len(result) + 1, dtype=np.int64))
    quantity = pd.to_numeric(result["Quantidade"], errors="raise")
    unit_cost = pd.to_numeric(result["Custo_Unitario"], errors="raise")
    list_price = pd.to_numeric(result["Preco_Unitario_Lista"], errors="raise")
    net = pd.to_numeric(result["Valor_Compra"], errors="raise")
    result["Valor_Bruto_Item"] = quantity * list_price
    result["Custo_Total_Item"] = quantity * unit_cost
    result["Valor_Desconto_Item"] = result["Valor_Bruto_Item"] - net
    result["Margem_Bruta_Item"] = net - result["Custo_Total_Item"]
    result["Margem_Percentual_Item"] = np.where(net.ne(0), result["Margem_Bruta_Item"] / net, np.nan)
    columns = [
        "Item_Pedido_SK", "Pedido_ID", "Cliente_SK", "Produto_SK", "Data_SK",
        "Horario_SK", "Pagamento_SK", "Campanha_SK", "Canal_Marketing_SK",
        "Evento_Externo_SK", "Bronze_Source_Row_Number", "Quantidade",
        "Custo_Unitario", "Preco_Unitario_Lista", "Desconto_Percentual",
        "Valor_Compra", "Valor_Bruto_Item", "Custo_Total_Item",
        "Valor_Desconto_Item", "Margem_Bruta_Item", "Margem_Percentual_Item",
    ]
    return result[columns]


def build_order_fact(mapped):
    campaign_text = normalize_text(mapped["Campanha"]).str.casefold()
    mapped = mapped.copy()
    mapped["Pedido_Com_Campanha"] = ~campaign_text.isin(["sem campanha", "nenhuma", "não", "nao"])
    # Mantém o nulo até a consolidação do pedido. O agregador `first` ignora
    # nulos e preserva um cupom válido caso ele exista em outra linha do pedido.
    mapped["Cupom_Status"] = normalize_text(mapped["Cupom_Utilizado"])
    aggregations = {
        "Cliente_SK": "first", "Data_SK": "first", "Horario_SK": "first",
        "Pagamento_SK": "first", "Campanha_SK": "first", "Canal_Marketing_SK": "first",
        "Evento_Externo_SK": "first", "Geografia_SK": "first",
        "Produto_SK": pd.Series.nunique, "Quantidade": "sum", "Valor_Compra": "sum",
        "Pontos_Fidelidade": first_value, "Data_Hora_Compra": "first",
        "Cupom_Status": "first", "Pedido_Com_Campanha": "first",
    }
    result = mapped.groupby("Pedido_ID", as_index=False).agg(aggregations)
    result["Cupom_Status"] = result["Cupom_Status"].fillna("Não informado")
    line_counts = mapped.groupby("Pedido_ID").size().rename("Quantidade_Linhas_Pedido")
    gross = (pd.to_numeric(mapped["Quantidade"]) * pd.to_numeric(mapped["Preco_Unitario_Lista"])).groupby(mapped["Pedido_ID"]).sum()
    cost = (pd.to_numeric(mapped["Quantidade"]) * pd.to_numeric(mapped["Custo_Unitario"])).groupby(mapped["Pedido_ID"]).sum()
    result = result.merge(line_counts, on="Pedido_ID", validate="1:1")
    result["Valor_Bruto_Pedido"] = result["Pedido_ID"].map(gross)
    result["Custo_Total_Pedido"] = result["Pedido_ID"].map(cost)
    result = result.rename(columns={
        "Produto_SK": "Quantidade_Produtos_Distintos", "Quantidade": "Quantidade_Total_Itens",
        "Valor_Compra": "Valor_Liquido_Pedido", "Pontos_Fidelidade": "Pontos_Fidelidade_Observados",
    })
    result["Valor_Desconto_Pedido"] = result["Valor_Bruto_Pedido"] - result["Valor_Liquido_Pedido"]
    result["Margem_Bruta_Pedido"] = result["Valor_Liquido_Pedido"] - result["Custo_Total_Pedido"]
    result["Margem_Percentual_Pedido"] = np.where(result["Valor_Liquido_Pedido"].ne(0), result["Margem_Bruta_Pedido"] / result["Valor_Liquido_Pedido"], np.nan)
    result["Desconto_Medio_Ponderado"] = np.where(result["Valor_Bruto_Pedido"].ne(0), result["Valor_Desconto_Pedido"] / result["Valor_Bruto_Pedido"], np.nan)
    result = result.sort_values("Pedido_ID", kind="stable").reset_index(drop=True)
    result.insert(0, "Pedido_SK", np.arange(1, len(result) + 1, dtype=np.int64))
    columns = [
        "Pedido_SK", "Pedido_ID", "Cliente_SK", "Data_SK", "Horario_SK",
        "Pagamento_SK", "Campanha_SK", "Canal_Marketing_SK", "Evento_Externo_SK",
        "Geografia_SK", "Quantidade_Linhas_Pedido", "Quantidade_Total_Itens",
        "Quantidade_Produtos_Distintos", "Valor_Bruto_Pedido", "Valor_Desconto_Pedido",
        "Valor_Liquido_Pedido", "Custo_Total_Pedido", "Margem_Bruta_Pedido",
        "Margem_Percentual_Pedido", "Desconto_Medio_Ponderado", "Cupom_Status",
        "Pedido_Com_Campanha", "Pontos_Fidelidade_Observados", "Data_Hora_Compra",
    ]
    return result[columns]


def build_delivery_fact(mapped, order_fact):
    source = mapped.groupby("Pedido_ID", as_index=False).agg({
        "Cliente_SK": "first", "Data_SK": "first", "Geografia_SK": "first",
        "Transportadora_SK": "first", "Frete": first_value, "Frete_Gratis": first_value,
        "Prazo_Entrega_Prometido": first_value, "Prazo_Entrega_Real": first_value,
        "Atraso_Entrega": first_value, "Avaliacao_Cliente": first_value,
    })
    source = source.rename(columns={
        "Data_SK": "Data_Pedido_SK", "Prazo_Entrega_Prometido": "Prazo_Entrega_Prometido_Dias",
        "Prazo_Entrega_Real": "Prazo_Entrega_Real_Dias", "Atraso_Entrega": "Atraso_Entrega_Informado",
    })
    source["Diferenca_Prazo_Dias"] = pd.to_numeric(source["Prazo_Entrega_Real_Dias"]) - pd.to_numeric(source["Prazo_Entrega_Prometido_Dias"])
    source["Dias_Atraso"] = source["Diferenca_Prazo_Dias"].clip(lower=0)
    source["Dias_Antecipacao"] = (-source["Diferenca_Prazo_Dias"]).clip(lower=0)
    source["Status_Entrega_Calculado"] = np.select(
        [source["Diferenca_Prazo_Dias"].gt(0), source["Diferenca_Prazo_Dias"].lt(0)],
        ["Atrasada", "Antecipada"], default="No prazo",
    )
    context = order_fact[["Pedido_ID", "Valor_Liquido_Pedido", "Quantidade_Total_Itens"]]
    source = source.merge(context, on="Pedido_ID", how="left", validate="1:1")
    source = source.sort_values("Pedido_ID", kind="stable").reset_index(drop=True)
    source.insert(0, "Entrega_SK", np.arange(1, len(source) + 1, dtype=np.int64))
    columns = [
        "Entrega_SK", "Pedido_ID", "Cliente_SK", "Data_Pedido_SK", "Geografia_SK",
        "Transportadora_SK", "Frete", "Frete_Gratis", "Prazo_Entrega_Prometido_Dias",
        "Prazo_Entrega_Real_Dias", "Diferenca_Prazo_Dias", "Dias_Atraso",
        "Dias_Antecipacao", "Status_Entrega_Calculado", "Atraso_Entrega_Informado",
        "Avaliacao_Cliente", "Valor_Liquido_Pedido", "Quantidade_Total_Itens",
    ]
    return source[columns]


def build_client_month_fact(mapped, order_fact):
    order_dates = order_fact[["Pedido_ID", "Cliente_SK", "Data_Hora_Compra", "Valor_Liquido_Pedido", "Custo_Total_Pedido", "Margem_Bruta_Pedido", "Quantidade_Total_Itens", "Pontos_Fidelidade_Observados"]].copy()
    order_dates["Mes"] = order_dates["Data_Hora_Compra"].dt.to_period("M")
    product_context = mapped[["Pedido_ID", "Produto_SK", "Categoria_Item"]].drop_duplicates()
    product_context = product_context.merge(order_fact[["Pedido_ID", "Cliente_SK", "Data_Hora_Compra"]], on="Pedido_ID", validate="m:1")
    product_context["Mes"] = product_context["Data_Hora_Compra"].dt.to_period("M")
    product_month = product_context.groupby(["Cliente_SK", "Mes"]).agg(
        Produtos_Distintos_Mes=("Produto_SK", "nunique"),
        Categorias_Distintas_Mes=("Categoria_Item", "nunique"),
    ).reset_index()
    monthly = order_dates.groupby(["Cliente_SK", "Mes"]).agg(
        Pedidos_Mes=("Pedido_ID", "nunique"), Itens_Mes=("Quantidade_Total_Itens", "sum"),
        Receita_Liquida_Mes=("Valor_Liquido_Pedido", "sum"),
        Custo_Total_Mes=("Custo_Total_Pedido", "sum"),
        Margem_Bruta_Mes=("Margem_Bruta_Pedido", "sum"),
    ).reset_index().merge(product_month, on=["Cliente_SK", "Mes"], how="left", validate="1:1")

    grid = []
    for client_sk, group in order_dates.groupby("Cliente_SK"):
        first_month = group["Mes"].min()
        for month in pd.period_range(first_month, pd.Period("2025-12", freq="M"), freq="M"):
            grid.append((client_sk, month))
    result = pd.DataFrame(grid, columns=["Cliente_SK", "Mes"])
    result = result.merge(monthly, on=["Cliente_SK", "Mes"], how="left", validate="1:1")
    zero_columns = ["Pedidos_Mes", "Itens_Mes", "Produtos_Distintos_Mes", "Categorias_Distintas_Mes", "Receita_Liquida_Mes", "Custo_Total_Mes", "Margem_Bruta_Mes"]
    result[zero_columns] = result[zero_columns].fillna(0)
    result["Ticket_Medio_Calculado"] = np.where(result["Pedidos_Mes"].gt(0), result["Receita_Liquida_Mes"] / result["Pedidos_Mes"], np.nan)
    result["Data_Corte"] = result["Mes"].dt.to_timestamp(how="end").dt.normalize()
    result["Data_Corte_SK"] = result["Data_Corte"].dt.strftime("%Y%m%d").astype("int32")

    recency, frequency, interval, points = [], [], [], []
    for row in result.itertuples(index=False):
        history = order_dates[(order_dates["Cliente_SK"] == row.Cliente_SK) & (order_dates["Data_Hora_Compra"].dt.normalize() <= row.Data_Corte)].sort_values("Data_Hora_Compra")
        recency.append((row.Data_Corte - history["Data_Hora_Compra"].max().normalize()).days)
        frequency.append(history["Pedido_ID"].nunique())
        unique_dates = history["Data_Hora_Compra"].dt.normalize().drop_duplicates().sort_values()
        interval.append(unique_dates.diff().dt.days.mean() if len(unique_dates) > 1 else np.nan)
        observed = history["Pontos_Fidelidade_Observados"].dropna()
        points.append(observed.iloc[-1] if not observed.empty else pd.NA)
    result["Recencia_Dias"] = recency
    result["Frequencia_Acumulada"] = frequency
    result["Intervalo_Medio_Calculado"] = interval
    result["Pontos_Fidelidade_Ultimo_Observado"] = points
    result["Status_Cliente_Calculado"] = pd.Series(pd.NA, index=result.index, dtype="string")
    result = result.sort_values(["Cliente_SK", "Data_Corte_SK"], kind="stable").reset_index(drop=True)
    result.insert(0, "Cliente_Mes_SK", np.arange(1, len(result) + 1, dtype=np.int64))
    columns = [
        "Cliente_Mes_SK", "Cliente_SK", "Data_Corte_SK", "Pedidos_Mes", "Itens_Mes",
        "Produtos_Distintos_Mes", "Categorias_Distintas_Mes", "Receita_Liquida_Mes",
        "Custo_Total_Mes", "Margem_Bruta_Mes", "Ticket_Medio_Calculado", "Recencia_Dias",
        "Frequencia_Acumulada", "Intervalo_Medio_Calculado",
        "Pontos_Fidelidade_Ultimo_Observado", "Status_Cliente_Calculado",
    ]
    return result[columns]


def validate_columns(tables, specification):
    failures = []
    for name, frame in tables.items():
        expected = specification.loc[specification["tabela"].eq(name), "coluna_gold"].tolist()
        missing = [c for c in expected if c not in frame.columns]
        extra = [c for c in frame.columns if c not in expected]
        if missing or extra:
            failures.append({"tabela": name, "colunas_ausentes": str(missing), "colunas_extras": str(extra)})
        else:
            tables[name] = frame[expected]
    if failures:
        raise ValueError(f"Divergência entre materialização e contrato: {failures}")


def reconciliation_report(tables, mapped, client_conflicts, order_conflicts, calendar):
    item = tables["fato_item_pedido"]
    order = tables["fato_pedido"]
    client_month = tables["fato_cliente_mes"]
    fk_pairs = [
        ("fato_item_pedido", "Cliente_SK", "dim_cliente", "Cliente_SK"),
        ("fato_item_pedido", "Produto_SK", "dim_produto", "Produto_SK"),
        ("fato_item_pedido", "Data_SK", "dim_data", "Data_SK"),
        ("fato_item_pedido", "Horario_SK", "dim_horario", "Horario_SK"),
        ("fato_item_pedido", "Pagamento_SK", "dim_pagamento", "Pagamento_SK"),
        ("fato_item_pedido", "Campanha_SK", "dim_campanha", "Campanha_SK"),
        ("fato_item_pedido", "Canal_Marketing_SK", "dim_canal_marketing", "Canal_Marketing_SK"),
        ("fato_item_pedido", "Evento_Externo_SK", "dim_evento_externo", "Evento_Externo_SK"),
        ("fato_pedido", "Cliente_SK", "dim_cliente", "Cliente_SK"),
        ("fato_pedido", "Data_SK", "dim_data", "Data_SK"),
        ("fato_pedido", "Horario_SK", "dim_horario", "Horario_SK"),
        ("fato_pedido", "Geografia_SK", "dim_geografia", "Geografia_SK"),
        ("fato_pedido", "Pagamento_SK", "dim_pagamento", "Pagamento_SK"),
        ("fato_pedido", "Campanha_SK", "dim_campanha", "Campanha_SK"),
        ("fato_pedido", "Canal_Marketing_SK", "dim_canal_marketing", "Canal_Marketing_SK"),
        ("fato_pedido", "Evento_Externo_SK", "dim_evento_externo", "Evento_Externo_SK"),
        ("fato_entrega", "Cliente_SK", "dim_cliente", "Cliente_SK"),
        ("fato_entrega", "Data_Pedido_SK", "dim_data", "Data_SK"),
        ("fato_entrega", "Geografia_SK", "dim_geografia", "Geografia_SK"),
        ("fato_entrega", "Transportadora_SK", "dim_transportadora", "Transportadora_SK"),
        ("fato_cliente_mes", "Cliente_SK", "dim_cliente", "Cliente_SK"),
        ("fato_cliente_mes", "Data_Corte_SK", "dim_data", "Data_SK"),
    ]
    orphan_count = sum((~tables[fact][fk].isin(tables[dim][pk])).sum() for fact, fk, dim, pk in fk_pairs)
    continuity_failures = 0
    for _, group in client_month.groupby("Cliente_SK"):
        dates = pd.to_datetime(group["Data_Corte_SK"].astype(str), format="%Y%m%d").dt.to_period("M")
        expected = pd.period_range(dates.min(), pd.Period("2025-12", freq="M"), freq="M")
        continuity_failures += int(list(dates.sort_values()) != list(expected))
    # Reconcilia os pedidos integralmente sem cupom entre a Silver e a Gold.
    # A quantidade não é fixa: ausência de cupom é uma condição legítima e
    # deve acompanhar o conteúdo efetivo da fonte.
    coupon_source = normalize_text(mapped["Cupom_Utilizado"])
    expected_coupon_null_orders = int(
        coupon_source.isna().groupby(mapped["Pedido_ID"]).all().sum()
    )
    observed_coupon_null_orders = int(order["Cupom_Status"].eq("Não informado").sum())
    checks = [
        ("REC001", len(item), EXPECTED_ROWS, len(item) == EXPECTED_ROWS),
        ("REC002", item["Bronze_Source_Row_Number"].nunique(), EXPECTED_ROWS, item["Bronze_Source_Row_Number"].nunique() == EXPECTED_ROWS),
        ("REC003", len(order), EXPECTED_ORDERS, len(order) == order["Pedido_ID"].nunique() == EXPECTED_ORDERS),
        ("REC004", tables["dim_cliente"]["Cliente_ID"].nunique(), EXPECTED_CLIENTS, tables["dim_cliente"]["Cliente_ID"].nunique() == EXPECTED_CLIENTS),
        ("REC005", tables["dim_produto"]["Id_Item"].nunique(), EXPECTED_PRODUCTS, tables["dim_produto"]["Id_Item"].nunique() == EXPECTED_PRODUCTS),
        ("REC006", len(tables["fato_entrega"]), EXPECTED_ORDERS, len(tables["fato_entrega"]) == tables["fato_entrega"]["Pedido_ID"].nunique() == EXPECTED_ORDERS),
        ("REC007", float(item["Valor_Compra"].sum() - order["Valor_Liquido_Pedido"].sum()), 0, abs(item["Valor_Compra"].sum() - order["Valor_Liquido_Pedido"].sum()) <= TOLERANCE),
        ("REC008", float(item["Quantidade"].sum() - order["Quantidade_Total_Itens"].sum()), 0, item["Quantidade"].sum() == order["Quantidade_Total_Itens"].sum()),
        ("REC009", int(orphan_count), 0, orphan_count == 0),
        ("REC010", len(tables["dim_data"]), 365, len(tables["dim_data"]) == 365),
        ("REC011", len(tables["dim_horario"]), 1440, len(tables["dim_horario"]) == 1440),
        ("REC012", observed_coupon_null_orders, expected_coupon_null_orders, observed_coupon_null_orders == expected_coupon_null_orders),
        ("REC013", int((order_conflicts["atributo"] == "Pontos_Fidelidade").sum()) if not order_conflicts.empty else 0, 0, order_conflicts.empty or not order_conflicts["atributo"].eq("Pontos_Fidelidade").any()),
        ("REC014", continuity_failures, 0, continuity_failures == 0),
        ("REC015", int(tables["dim_data"]["Nome_Dia_Especial"].notna().sum()), len(calendar), tables["dim_data"]["Nome_Dia_Especial"].notna().sum() == len(calendar)),
        ("REC016", len(client_conflicts), 0, client_conflicts.empty),
    ]
    return pd.DataFrame([
        {"regra_id": rid, "observado": observed, "esperado": expected, "status": "aprovado" if passed else "reprovado"}
        for rid, observed, expected, passed in checks
    ])


def save_parquets(tables):
    GOLD_DIR.mkdir(parents=True, exist_ok=True)
    paths = {}
    for name, frame in tables.items():
        path = GOLD_DIR / f"{name}.parquet"
        temporary = GOLD_DIR / f".{name}.{uuid.uuid4().hex}.tmp.parquet"
        frame.to_parquet(temporary, index=False, engine="pyarrow")
        temporary.replace(path)
        paths[name] = path
    return paths


def save_report(sheets):
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    temporary = REPORT_DIR / f".{REPORT_PATH.stem}.{uuid.uuid4().hex}.tmp.xlsx"
    try:
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            for name, frame in sheets:
                frame.to_excel(writer, sheet_name=name[:31], index=False)
                ws = writer.book[name[:31]]
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="1F4E78")
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                for column in ws.columns:
                    width = min(max(len(str(c.value)) if c.value is not None else 0 for c in column) + 2, 55)
                    ws.column_dimensions[column[0].column_letter].width = max(width, 12)
        temporary.replace(REPORT_PATH)
    except PermissionError as error:
        temporary.unlink(missing_ok=True)
        raise PermissionError(f"Feche '{REPORT_PATH.name}' no Excel e aguarde o OneDrive.") from error


def main():
    print("Iniciando materialização do modelo dimensional Gold...")
    if not SILVER_PATH.is_file():
        raise FileNotFoundError(f"Silver ausente: {SILVER_PATH}")
    if not SPEC_PATH.is_file():
        raise FileNotFoundError(f"Especificação da Etapa 11 ausente: {SPEC_PATH}")

    silver = pd.read_parquet(SILVER_PATH, engine="pyarrow")
    if "_bronze_source_row_number" in silver.columns:
        silver = silver.rename(
            columns={"_bronze_source_row_number": "Bronze_Source_Row_Number"}
        )
    specification = pd.read_excel(SPEC_PATH, sheet_name="colunas_modelo")
    calendar = pd.read_excel(SPEC_PATH, sheet_name="calendario_2025")
    validate_inputs(silver, specification)

    client_conflicts = conflict_report(silver, "Cliente_ID", CLIENT_ATTRIBUTES, "dim_cliente")
    order_conflicts = conflict_report(silver, "Pedido_ID", ORDER_ATTRIBUTES, "fato_pedido/fato_entrega")
    blocking_order_conflicts = order_conflicts[order_conflicts["atributo"].isin(ORDER_ATTRIBUTES)]
    if not client_conflicts.empty:
        raise ValueError(f"REC016 reprovada: {len(client_conflicts)} conflitos de atributos por cliente.")
    if not blocking_order_conflicts.empty:
        raise ValueError(f"Existem {len(blocking_order_conflicts)} conflitos dentro de pedidos.")

    run_time = utc_now()
    load_time = run_time.replace(tzinfo=None)
    dimensions = {
        "dim_data": build_date_dimension(calendar),
        "dim_horario": build_time_dimension(),
        "dim_geografia": build_geography_dimension(silver),
        "dim_pagamento": deterministic_dimension(silver["Forma_Pagamento"], "Pagamento_SK", "Forma_Pagamento"),
        "dim_campanha": deterministic_dimension(silver["Campanha"], "Campanha_SK", "Campanha"),
        "dim_canal_marketing": deterministic_dimension(silver["Canal_Marketing"], "Canal_Marketing_SK", "Canal_Marketing"),
        "dim_evento_externo": deterministic_dimension(silver["Evento_Externo"], "Evento_Externo_SK", "Evento_Externo"),
        "dim_transportadora": deterministic_dimension(silver["Transportadora"], "Transportadora_SK", "Transportadora"),
    }
    dimensions["dim_cliente"] = build_client_dimension(silver, dimensions["dim_geografia"], load_time)
    dimensions["dim_produto"] = build_product_dimension(silver)
    mapped = map_foreign_keys(silver, dimensions)
    facts = {}
    facts["fato_item_pedido"] = build_item_fact(mapped)
    facts["fato_pedido"] = build_order_fact(mapped)
    facts["fato_entrega"] = build_delivery_fact(mapped, facts["fato_pedido"])
    facts["fato_cliente_mes"] = build_client_month_fact(mapped, facts["fato_pedido"])
    tables = {**dimensions, **facts}
    validate_columns(tables, specification)
    reconciliation = reconciliation_report(tables, mapped, client_conflicts, order_conflicts, calendar)
    if reconciliation["status"].eq("reprovado").any():
        failures = reconciliation.loc[reconciliation["status"].eq("reprovado"), "regra_id"].tolist()
        raise ValueError(f"Reconciliações Gold reprovadas: {failures}")

    paths = save_parquets(tables)
    inventory = pd.DataFrame([
        {"tabela": name, "linhas": len(tables[name]), "colunas": len(tables[name].columns), "arquivo": str(path.relative_to(PROJECT_ROOT)), "sha256": sha256_file(path)}
        for name, path in sorted(paths.items())
    ])
    summary = pd.DataFrame([{
        "executado_em_utc": run_time.isoformat(), "etapa": "12_materializacao_modelo_dimensional_gold",
        "sha256_silver": sha256_file(SILVER_PATH), "sha256_especificacao": sha256_file(SPEC_PATH),
        "tabelas_materializadas": len(tables), "dimensoes": len(dimensions), "fatos": len(facts),
        "regras_aprovadas": int(reconciliation["status"].eq("aprovado").sum()),
        "resultado": "APROVADO", "dados_silver_modificados": False,
    }])
    samples = []
    for name, frame in sorted(tables.items()):
        sample = frame.head(20).copy()
        sample.insert(0, "tabela_origem", name)
        samples.append(sample)
    save_report([
        ("resumo_execucao", summary), ("inventario_gold", inventory),
        ("reconciliacoes", reconciliation), ("conflitos_cliente", client_conflicts),
        ("conflitos_pedido", order_conflicts), ("amostras_gold", pd.concat(samples, ignore_index=True, sort=False)),
    ])
    manifest = {
        "run_utc": run_time.isoformat(), "status": "success",
        "inputs": {
            "silver": {"path": str(SILVER_PATH.relative_to(PROJECT_ROOT)), "sha256": sha256_file(SILVER_PATH)},
            "specification": {"path": str(SPEC_PATH.relative_to(PROJECT_ROOT)), "sha256": sha256_file(SPEC_PATH)},
        },
        "outputs": {
            name: {"path": str(path.relative_to(PROJECT_ROOT)), "sha256": sha256_file(path), "rows": len(tables[name]), "columns": len(tables[name].columns)}
            for name, path in sorted(paths.items())
        },
        "report": {"path": str(REPORT_PATH.relative_to(PROJECT_ROOT)), "sha256": sha256_file(REPORT_PATH)},
        "reconciliations": reconciliation.to_dict(orient="records"),
    }
    MANIFEST_DIR.mkdir(parents=True, exist_ok=True)
    temporary_manifest = MANIFEST_DIR / f".{MANIFEST_PATH.stem}.{uuid.uuid4().hex}.tmp.json"
    temporary_manifest.write_text(json.dumps(manifest, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    temporary_manifest.replace(MANIFEST_PATH)
    print(f"[OK] 14 tabelas Gold materializadas em: {GOLD_DIR}")
    print(f"[OK] Relatório: {REPORT_PATH}")
    print(f"[OK] Manifesto: {MANIFEST_PATH}")
    print("RESULTADO_FINAL: APROVADO")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print("RESULTADO_FINAL: REPROVADO")
        print(f"ERRO: {type(error).__name__}: {error}")
        sys.exit(1)
