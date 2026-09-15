import hashlib
import json
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


PROJECT_ROOT = Path(__file__).resolve().parents[3]
SILVER_PATH = (
    PROJECT_ROOT / "data" / "silver" / "food_commerce"
    / "silver_transactions_v3.parquet"
)
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "modeling" / "gold_v3"
OUTPUT_PATH = OUTPUT_DIR / "step_11_especificacao_modelo_dimensional.xlsx"
MANIFEST_DIR = PROJECT_ROOT / "outputs" / "manifests" / "gold_v3"
MANIFEST_PATH = MANIFEST_DIR / "step_11_especificacao_modelo_dimensional.json"

EXPECTED_ROWS = 6495
EXPECTED_ORDERS = 2789
EXPECTED_CLIENTS = 330
EXPECTED_PRODUCTS = 36

BUSINESS_COLUMNS = [
    "Pedido_ID", "Cliente_ID", "Idade", "Gênero", "Cidade", "Estado",
    "Estado_Civil", "Escolaridade", "Situacao_Profissional", "Profissão",
    "Tempo_Experiencia_Anos", "Renda_Mensal_Cliente", "Ciclo_Vida_Cliente",
    "Intervalo_Medio_Compras_Dias", "Categoria_Favorita",
    "Forma_Pagamento_Preferida", "Dia_Semana_Preferido",
    "Ticket_Medio_Cliente", "Id_Item", "Categoria_Item", "Produto", "Medida",
    "Quantidade", "Custo_Unitario", "Preco_Unitario_Lista", "Valor_Compra",
    "Forma_Pagamento", "Data_Compra", "Horario_Compra", "Mes_Ref", "Campanha",
    "Evento_Externo", "Canal_Marketing", "Cupom_Utilizado",
    "Desconto_Percentual", "Frete", "Frete_Gratis", "Transportadora",
    "Prazo_Entrega_Prometido", "Prazo_Entrega_Real", "Atraso_Entrega",
    "Avaliacao_Cliente", "Pontos_Fidelidade", "Data_Hora_Compra",
]


def sha256_file(path):
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def add_columns(rows, table, grain, area, definitions):
    for name, dtype, role, source, nullable, rule in definitions:
        rows.append({
            "tabela": table,
            "granularidade": grain,
            "areas_consumidoras": area,
            "coluna_gold": name,
            "tipo_sugerido": dtype,
            "papel": role,
            "origem_silver": source,
            "aceita_nulo": nullable,
            "regra_ou_interpretacao": rule,
        })


def build_tables():
    return pd.DataFrame([
        ["dim_cliente", "dimensao", "um cliente atual (SCD Tipo 1)", "Cliente_SK", "Cliente_ID", "330", "todas"],
        ["dim_produto", "dimensao", "um produto", "Produto_SK", "Id_Item", "36", "comercial;financeira;marketing"],
        ["dim_data", "dimensao", "um dia do calendario", "Data_SK", "Data_Completa", "365 para 2025", "todas"],
        ["dim_horario", "dimensao", "um minuto do dia", "Horario_SK", "Horario", "1440", "comercial;marketing"],
        ["dim_geografia", "dimensao", "uma combinacao cidade-estado", "Geografia_SK", "Cidade+Estado", "conforme Silver", "todas"],
        ["dim_pagamento", "dimensao", "uma forma de pagamento", "Pagamento_SK", "Forma_Pagamento", "conforme Silver", "comercial;financeira"],
        ["dim_campanha", "dimensao", "uma campanha", "Campanha_SK", "Campanha", "conforme Silver", "comercial;marketing"],
        ["dim_canal_marketing", "dimensao", "um canal", "Canal_Marketing_SK", "Canal_Marketing", "conforme Silver", "comercial;marketing"],
        ["dim_evento_externo", "dimensao", "um evento externo", "Evento_Externo_SK", "Evento_Externo", "conforme Silver", "comercial;marketing"],
        ["dim_transportadora", "dimensao", "uma transportadora", "Transportadora_SK", "Transportadora", "conforme Silver", "logistica"],
        ["fato_item_pedido", "fato_transacional", "uma linha de produto no pedido", "Item_Pedido_SK", "Pedido_ID+Bronze_Source_Row_Number", "6495", "comercial;financeira;marketing"],
        ["fato_pedido", "fato_transacional", "um pedido", "Pedido_SK", "Pedido_ID", "2789", "comercial;financeira;marketing"],
        ["fato_entrega", "fato_transacional", "uma entrega por pedido", "Entrega_SK", "Pedido_ID", "2789", "logistica;comercial"],
        ["fato_cliente_mes", "fato_snapshot_periodico", "um cliente por fechamento mensal desde a primeira compra observada", "Cliente_Mes_SK", "Cliente_SK+Data_Corte_SK", "derivada", "todas"],
    ], columns=["tabela", "tipo_tabela", "granularidade", "chave_primaria", "chave_natural", "quantidade_esperada", "areas_consumidoras"])


def build_columns():
    rows = []
    add_columns(rows, "dim_cliente", "um cliente atual (SCD Tipo 1)", "todas", [
        ("Cliente_SK", "int64", "PK", "gerada", "nao", "chave substituta deterministica"),
        ("Cliente_ID", "string", "NK", "Cliente_ID", "nao", "identificador natural preservado"),
        ("Geografia_SK", "int64", "FK", "Cidade+Estado", "nao", "referencia dim_geografia"),
        ("Idade", "int64", "atributo", "Idade", "nao", "perfil demografico"),
        ("Genero", "string", "atributo", "Gênero", "nao", "descricao normalizada"),
        ("Estado_Civil", "string", "atributo", "Estado_Civil", "nao", "valor Silver"),
        ("Escolaridade", "string", "atributo", "Escolaridade", "nao", "valor Silver"),
        ("Situacao_Profissional", "string", "atributo", "Situacao_Profissional", "nao", "valor Silver"),
        ("Profissao", "string", "atributo", "Profissão", "nao", "valor Silver"),
        ("Tempo_Experiencia_Anos", "int64", "atributo", "Tempo_Experiencia_Anos", "nao", "valor Silver"),
        ("Renda_Mensal_Cliente", "float64", "atributo", "Renda_Mensal_Cliente", "sim", "um nulo documentado deve ser preservado"),
        ("Ciclo_Vida_Informado", "string", "atributo_referencia", "Ciclo_Vida_Cliente", "nao", "nao substituir pelo status recalculado"),
        ("Intervalo_Medio_Informado", "float64", "atributo_referencia", "Intervalo_Medio_Compras_Dias", "nao", "comparar com metrica recalculada"),
        ("Categoria_Favorita_Informada", "string", "atributo_referencia", "Categoria_Favorita", "nao", "comparar com preferencia observada"),
        ("Pagamento_Preferido_Informado", "string", "atributo_referencia", "Forma_Pagamento_Preferida", "nao", "comparar com comportamento observado"),
        ("Dia_Semana_Preferido_Informado", "string", "atributo_referencia", "Dia_Semana_Preferido", "nao", "comparar com comportamento observado"),
        ("Ticket_Medio_Informado", "float64", "atributo_referencia", "Ticket_Medio_Cliente", "nao", "comparar com ticket recalculado"),
        ("Data_Carga", "datetime", "auditoria", "gerada", "nao", "momento da materializacao"),
        ("Data_Atualizacao", "datetime", "auditoria", "gerada", "nao", "igual a Data_Carga na carga inicial"),
        ("Hash_Atributos", "string", "auditoria", "gerada", "nao", "detecta alteracao de atributos"),
    ])
    add_columns(rows, "dim_produto", "um produto", "comercial;financeira;marketing", [
        ("Produto_SK", "int64", "PK", "gerada", "nao", "chave substituta deterministica"),
        ("Id_Item", "string", "NK", "Id_Item", "nao", "identificador natural"),
        ("Produto", "string", "atributo", "Produto", "nao", "nome canonico Silver"),
        ("Categoria_Item", "string", "atributo", "Categoria_Item", "nao", "categoria canonica Silver"),
        ("Medida", "string", "atributo", "Medida", "nao", "embalagem ou medida"),
        ("Hash_Atributos", "string", "auditoria", "gerada", "nao", "detecta alteracoes"),
        ("Registro_Ativo", "bool", "controle", "gerada", "nao", "verdadeiro na carga inicial"),
    ])
    add_columns(rows, "dim_data", "um dia do calendario", "todas", [
        ("Data_SK", "int32", "PK", "gerada", "nao", "AAAAMMDD"),
        ("Data_Completa", "date", "NK", "calendario", "nao", "2025-01-01 a 2025-12-31"),
        ("Ano", "int16", "atributo", "derivada", "nao", "ano calendario"),
        ("Semestre", "int8", "atributo", "derivada", "nao", "1 ou 2"),
        ("Trimestre", "int8", "atributo", "derivada", "nao", "1 a 4"),
        ("Numero_Mes", "int8", "atributo", "derivada", "nao", "1 a 12"),
        ("Nome_Mes", "string", "atributo", "derivada", "nao", "nome em portugues"),
        ("Ano_Mes", "int32", "atributo", "derivada", "nao", "AAAAMM"),
        ("Dia_Mes", "int8", "atributo", "derivada", "nao", "1 a 31"),
        ("Dia_Ano", "int16", "atributo", "derivada", "nao", "1 a 365/366"),
        ("Semana_Ano", "int8", "atributo", "derivada", "nao", "padrao ISO"),
        ("Numero_Dia_Semana", "int8", "atributo", "derivada", "nao", "segunda=1"),
        ("Nome_Dia_Semana", "string", "atributo", "derivada", "nao", "nome em portugues"),
        ("Final_Semana", "bool", "atributo", "derivada", "nao", "sabado ou domingo"),
        ("Feriado_Nacional", "bool", "atributo_externo", "calendario_oficial_2025", "nao", "true apenas para feriado nacional oficial"),
        ("Ponto_Facultativo_Federal", "bool", "atributo_externo", "calendario_oficial_2025", "nao", "nao equivale a feriado nacional"),
        ("Nome_Dia_Especial", "string", "atributo_externo", "calendario_oficial_2025", "sim", "nome oficial do feriado ou ponto facultativo"),
        ("Tipo_Dia_Especial", "string", "atributo_externo", "calendario_oficial_2025", "sim", "feriado_nacional;ponto_facultativo_federal"),
        ("Escopo_Dia_Especial", "string", "atributo_externo", "calendario_oficial_2025", "sim", "nacional ou administracao_publica_federal"),
        ("Periodo_Dia_Especial", "string", "atributo_externo", "calendario_oficial_2025", "sim", "dia_inteiro;ate_14h;apos_13h"),
        ("Fonte_Calendario", "string", "linhagem_externa", "calendario_oficial_2025", "sim", "portaria e URL de origem"),
    ])
    add_columns(rows, "dim_horario", "um minuto do dia", "comercial;marketing", [
        ("Horario_SK", "int16", "PK", "gerada", "nao", "0 a 1439"),
        ("Horario", "time", "NK", "calendario", "nao", "HH:MM"),
        ("Hora", "int8", "atributo", "derivada", "nao", "0 a 23"),
        ("Minuto", "int8", "atributo", "derivada", "nao", "0 a 59"),
        ("Faixa_Horaria", "string", "atributo", "derivada", "nao", "madrugada;manha;tarde;noite"),
    ])
    for table, sk, natural, area in [
        ("dim_geografia", "Geografia_SK", "Cidade+Estado", "todas"),
        ("dim_pagamento", "Pagamento_SK", "Forma_Pagamento", "comercial;financeira"),
        ("dim_campanha", "Campanha_SK", "Campanha", "comercial;marketing"),
        ("dim_canal_marketing", "Canal_Marketing_SK", "Canal_Marketing", "comercial;marketing"),
        ("dim_evento_externo", "Evento_Externo_SK", "Evento_Externo", "comercial;marketing"),
        ("dim_transportadora", "Transportadora_SK", "Transportadora", "logistica"),
    ]:
        attributes = [(sk, "int64", "PK", "gerada", "nao", "chave substituta deterministica")]
        if table == "dim_geografia":
            attributes += [
                ("Cidade", "string", "NK_parcial", "Cidade", "nao", "cidade Silver"),
                ("Estado", "string", "NK_parcial", "Estado", "nao", "UF Silver"),
                ("Regiao", "string", "atributo", "derivada_UF", "nao", "macroregiao brasileira"),
                ("Pais", "string", "atributo", "constante", "nao", "Brasil"),
            ]
        else:
            attributes.append((natural, "string", "NK", natural, "nao", "categoria valida da Silver"))
        add_columns(rows, table, "uma categoria de referencia", area, attributes)

    fact_keys = [
        ("Item_Pedido_SK", "int64", "PK", "gerada", "nao", "chave da linha Gold"),
        ("Pedido_ID", "string", "dimensao_degenerada", "Pedido_ID", "nao", "preserva pedido sem dimensao propria"),
        ("Cliente_SK", "int64", "FK", "Cliente_ID", "nao", "referencia dim_cliente"),
        ("Produto_SK", "int64", "FK", "Id_Item", "nao", "referencia dim_produto"),
        ("Data_SK", "int32", "FK", "Data_Compra", "nao", "referencia dim_data"),
        ("Horario_SK", "int16", "FK", "Horario_Compra", "nao", "referencia dim_horario"),
        ("Pagamento_SK", "int64", "FK", "Forma_Pagamento", "nao", "referencia dim_pagamento"),
        ("Campanha_SK", "int64", "FK", "Campanha", "nao", "referencia dim_campanha"),
        ("Canal_Marketing_SK", "int64", "FK", "Canal_Marketing", "nao", "referencia dim_canal_marketing"),
        ("Evento_Externo_SK", "int64", "FK", "Evento_Externo", "nao", "referencia dim_evento_externo"),
        ("Bronze_Source_Row_Number", "int64", "linhagem", "_bronze_source_row_number", "nao", "identificador unico da origem"),
        ("Quantidade", "int64", "medida_aditiva", "Quantidade", "nao", "unidades na linha"),
        ("Custo_Unitario", "float64", "medida_nao_aditiva", "Custo_Unitario", "nao", "nao somar isoladamente"),
        ("Preco_Unitario_Lista", "float64", "medida_nao_aditiva", "Preco_Unitario_Lista", "nao", "nao somar isoladamente"),
        ("Desconto_Percentual", "float64", "medida_nao_aditiva", "Desconto_Percentual", "nao", "calcular media ponderada"),
        ("Valor_Compra", "float64", "medida_aditiva", "Valor_Compra", "nao", "valor liquido observado"),
        ("Valor_Bruto_Item", "float64", "medida_aditiva", "calculada", "nao", "Quantidade*Preco_Unitario_Lista"),
        ("Custo_Total_Item", "float64", "medida_aditiva", "calculada", "nao", "Quantidade*Custo_Unitario"),
        ("Valor_Desconto_Item", "float64", "medida_aditiva", "calculada", "nao", "Valor_Bruto_Item-Valor_Compra"),
        ("Margem_Bruta_Item", "float64", "medida_aditiva", "calculada", "nao", "Valor_Compra-Custo_Total_Item"),
        ("Margem_Percentual_Item", "float64", "medida_nao_aditiva", "calculada", "sim", "Margem_Bruta_Item/Valor_Compra"),
    ]
    add_columns(rows, "fato_item_pedido", "uma linha de produto no pedido", "comercial;financeira;marketing", fact_keys)

    add_columns(rows, "fato_pedido", "um pedido", "comercial;financeira;marketing", [
        ("Pedido_SK", "int64", "PK", "gerada", "nao", "chave substituta"),
        ("Pedido_ID", "string", "NK", "Pedido_ID", "nao", "unico na fato"),
        ("Cliente_SK", "int64", "FK", "Cliente_ID", "nao", "dim_cliente"),
        ("Data_SK", "int32", "FK", "Data_Compra", "nao", "dim_data"),
        ("Horario_SK", "int16", "FK", "Horario_Compra", "nao", "dim_horario"),
        ("Pagamento_SK", "int64", "FK", "Forma_Pagamento", "nao", "dim_pagamento"),
        ("Campanha_SK", "int64", "FK", "Campanha", "nao", "dim_campanha"),
        ("Canal_Marketing_SK", "int64", "FK", "Canal_Marketing", "nao", "dim_canal_marketing"),
        ("Evento_Externo_SK", "int64", "FK", "Evento_Externo", "nao", "dim_evento_externo"),
        ("Geografia_SK", "int64", "FK", "Cidade+Estado", "nao", "dim_geografia"),
        ("Quantidade_Linhas_Pedido", "int64", "medida_aditiva", "calculada", "nao", "contagem de linhas"),
        ("Quantidade_Total_Itens", "int64", "medida_aditiva", "calculada", "nao", "soma Quantidade"),
        ("Quantidade_Produtos_Distintos", "int64", "medida", "calculada", "nao", "nunique Produto_SK"),
        ("Valor_Bruto_Pedido", "float64", "medida_aditiva", "calculada", "nao", "soma Valor_Bruto_Item"),
        ("Valor_Desconto_Pedido", "float64", "medida_aditiva", "calculada", "nao", "soma Valor_Desconto_Item"),
        ("Valor_Liquido_Pedido", "float64", "medida_aditiva", "calculada", "nao", "soma Valor_Compra"),
        ("Custo_Total_Pedido", "float64", "medida_aditiva", "calculada", "nao", "soma Custo_Total_Item"),
        ("Margem_Bruta_Pedido", "float64", "medida_aditiva", "calculada", "nao", "soma Margem_Bruta_Item"),
        ("Margem_Percentual_Pedido", "float64", "medida_nao_aditiva", "calculada", "sim", "Margem_Bruta_Pedido/Valor_Liquido_Pedido"),
        ("Desconto_Medio_Ponderado", "float64", "medida_nao_aditiva", "calculada", "sim", "Valor_Desconto_Pedido/Valor_Bruto_Pedido"),
        ("Cupom_Status", "string", "indicador", "Cupom_Utilizado", "nao", "Sim;Nao;Nao informado sem alterar Silver"),
        ("Pedido_Com_Campanha", "bool", "indicador", "Campanha", "nao", "Campanha diferente de Sem campanha"),
        ("Pontos_Fidelidade_Observados", "int64", "medida_snapshot", "Pontos_Fidelidade", "nao", "preservar por pedido; nao somar ate conhecer a regra"),
        ("Data_Hora_Compra", "datetime", "timestamp_evento", "Data_Hora_Compra", "nao", "preserva segundos do evento original"),
    ])
    add_columns(rows, "fato_entrega", "uma entrega por pedido", "logistica;comercial", [
        ("Entrega_SK", "int64", "PK", "gerada", "nao", "chave substituta"),
        ("Pedido_ID", "string", "dimensao_degenerada", "Pedido_ID", "nao", "reconciliacao"),
        ("Cliente_SK", "int64", "FK", "Cliente_ID", "nao", "dim_cliente"),
        ("Data_Pedido_SK", "int32", "FK", "Data_Compra", "nao", "dim_data"),
        ("Geografia_SK", "int64", "FK", "Cidade+Estado", "nao", "dim_geografia"),
        ("Transportadora_SK", "int64", "FK", "Transportadora", "nao", "dim_transportadora"),
        ("Frete", "float64", "medida_aditiva", "Frete", "nao", "somar somente nesta fato"),
        ("Frete_Gratis", "bool", "indicador", "Frete_Gratis", "nao", "Sim/Não"),
        ("Prazo_Entrega_Prometido_Dias", "int64", "medida", "Prazo_Entrega_Prometido", "nao", "dias"),
        ("Prazo_Entrega_Real_Dias", "int64", "medida", "Prazo_Entrega_Real", "nao", "dias"),
        ("Diferenca_Prazo_Dias", "int64", "medida", "calculada", "nao", "real-prometido"),
        ("Dias_Atraso", "int64", "medida", "calculada", "nao", "max(diferenca,0)"),
        ("Dias_Antecipacao", "int64", "medida", "calculada", "nao", "max(-diferenca,0)"),
        ("Status_Entrega_Calculado", "string", "indicador", "calculada", "nao", "Atrasada;No prazo;Antecipada"),
        ("Atraso_Entrega_Informado", "string", "atributo_referencia", "Atraso_Entrega", "nao", "comparar com status calculado"),
        ("Avaliacao_Cliente", "int64", "medida_ordem", "Avaliacao_Cliente", "nao", "1 a 5; usar mediana e distribuicao"),
        ("Valor_Liquido_Pedido", "float64", "medida_contexto", "fato_pedido", "nao", "nao somar oficialmente nesta fato"),
        ("Quantidade_Total_Itens", "int64", "medida_contexto", "fato_pedido", "nao", "nao somar oficialmente nesta fato"),
    ])
    add_columns(rows, "fato_cliente_mes", "um cliente por fechamento mensal desde a primeira compra observada", "todas", [
        ("Cliente_Mes_SK", "int64", "PK", "gerada", "nao", "chave substituta"),
        ("Cliente_SK", "int64", "FK", "dim_cliente", "nao", "cliente"),
        ("Data_Corte_SK", "int32", "FK", "dim_data", "nao", "ultimo dia do mes"),
        ("Pedidos_Mes", "int64", "medida_aditiva", "fato_pedido", "nao", "contagem"),
        ("Itens_Mes", "int64", "medida_aditiva", "fato_item_pedido", "nao", "soma"),
        ("Produtos_Distintos_Mes", "int64", "medida", "fato_item_pedido", "nao", "nunique"),
        ("Categorias_Distintas_Mes", "int64", "medida", "fato_item_pedido", "nao", "nunique"),
        ("Receita_Liquida_Mes", "float64", "medida_aditiva", "fato_pedido", "nao", "soma"),
        ("Custo_Total_Mes", "float64", "medida_aditiva", "fato_pedido", "nao", "soma"),
        ("Margem_Bruta_Mes", "float64", "medida_aditiva", "fato_pedido", "nao", "soma"),
        ("Ticket_Medio_Calculado", "float64", "medida_nao_aditiva", "calculada", "sim", "receita/pedidos"),
        ("Recencia_Dias", "int64", "medida_snapshot", "calculada", "nao", "data de corte-ultima compra"),
        ("Frequencia_Acumulada", "int64", "medida_snapshot", "calculada", "nao", "pedidos ate a data de corte"),
        ("Intervalo_Medio_Calculado", "float64", "medida_snapshot", "calculada", "sim", "media entre pedidos"),
        ("Pontos_Fidelidade_Ultimo_Observado", "int64", "medida_snapshot", "fato_pedido", "sim", "ultimo valor conhecido no fechamento; nao somar"),
        ("Status_Cliente_Calculado", "string", "classificacao", "regra_negocio", "sim", "definir depois da EDA temporal"),
    ])
    return pd.DataFrame(rows)


def build_bus_matrix():
    dimensions = ["dim_cliente", "dim_produto", "dim_data", "dim_horario", "dim_geografia", "dim_pagamento", "dim_campanha", "dim_canal_marketing", "dim_evento_externo", "dim_transportadora"]
    usage = {
        "fato_item_pedido": {"dim_cliente", "dim_produto", "dim_data", "dim_horario", "dim_pagamento", "dim_campanha", "dim_canal_marketing", "dim_evento_externo"},
        "fato_pedido": {"dim_cliente", "dim_data", "dim_horario", "dim_geografia", "dim_pagamento", "dim_campanha", "dim_canal_marketing", "dim_evento_externo"},
        "fato_entrega": {"dim_cliente", "dim_data", "dim_geografia", "dim_transportadora"},
        "fato_cliente_mes": {"dim_cliente", "dim_data"},
    }
    rows = []
    for fact, used in usage.items():
        row = {"tabela_fato": fact}
        row.update({dimension: "X" if dimension in used else "" for dimension in dimensions})
        rows.append(row)
    return pd.DataFrame(rows)


def build_relationships():
    rows = []
    mappings = {
        "fato_item_pedido": ["Cliente_SK", "Produto_SK", "Data_SK", "Horario_SK", "Pagamento_SK", "Campanha_SK", "Canal_Marketing_SK", "Evento_Externo_SK"],
        "fato_pedido": ["Cliente_SK", "Data_SK", "Horario_SK", "Geografia_SK", "Pagamento_SK", "Campanha_SK", "Canal_Marketing_SK", "Evento_Externo_SK"],
        "fato_entrega": ["Cliente_SK", "Data_Pedido_SK", "Geografia_SK", "Transportadora_SK"],
        "fato_cliente_mes": ["Cliente_SK", "Data_Corte_SK"],
    }
    target = {
        "Cliente_SK": "dim_cliente.Cliente_SK", "Produto_SK": "dim_produto.Produto_SK",
        "Data_SK": "dim_data.Data_SK", "Data_Pedido_SK": "dim_data.Data_SK",
        "Data_Corte_SK": "dim_data.Data_SK", "Horario_SK": "dim_horario.Horario_SK",
        "Geografia_SK": "dim_geografia.Geografia_SK", "Pagamento_SK": "dim_pagamento.Pagamento_SK",
        "Campanha_SK": "dim_campanha.Campanha_SK", "Canal_Marketing_SK": "dim_canal_marketing.Canal_Marketing_SK",
        "Evento_Externo_SK": "dim_evento_externo.Evento_Externo_SK",
        "Transportadora_SK": "dim_transportadora.Transportadora_SK",
    }
    for fact, keys in mappings.items():
        for key in keys:
            rows.append({"tabela_origem": fact, "coluna_fk": key, "referencia": target[key], "cardinalidade": "N:1", "integridade_obrigatoria": "sim"})
    return pd.DataFrame(rows)


def build_metrics():
    return pd.DataFrame([
        ["quantidade_vendida", "fato_item_pedido", "SUM(Quantidade)", "aditiva", "comercial"],
        ["receita_produto", "fato_item_pedido", "SUM(Valor_Compra)", "aditiva", "comercial;financeira"],
        ["margem_produto", "fato_item_pedido", "SUM(Margem_Bruta_Item)", "aditiva", "financeira"],
        ["pedidos", "fato_pedido", "COUNT(Pedido_SK)", "aditiva", "comercial"],
        ["ticket_medio", "fato_pedido", "SUM(Valor_Liquido_Pedido)/COUNT(Pedido_SK)", "nao_aditiva", "comercial;financeira"],
        ["receita_liquida", "fato_pedido", "SUM(Valor_Liquido_Pedido)", "aditiva", "financeira"],
        ["desconto_total", "fato_pedido", "SUM(Valor_Desconto_Pedido)", "aditiva", "financeira;marketing"],
        ["margem_bruta", "fato_pedido", "SUM(Margem_Bruta_Pedido)", "aditiva", "financeira"],
        ["margem_percentual", "fato_pedido", "SUM(Margem_Bruta_Pedido)/SUM(Valor_Liquido_Pedido)", "nao_aditiva", "financeira"],
        ["frete_total", "fato_entrega", "SUM(Frete)", "aditiva", "logistica"],
        ["taxa_atraso", "fato_entrega", "AVG(indicador Entrega_Atrasada)", "nao_aditiva", "logistica"],
        ["prazo_medio_real", "fato_entrega", "AVG(Prazo_Entrega_Real_Dias)", "nao_aditiva", "logistica"],
        ["avaliacao_mediana", "fato_entrega", "MEDIAN(Avaliacao_Cliente)", "ordinal", "logistica"],
        ["clientes_ativos_mes", "fato_cliente_mes", "COUNT(Cliente_SK com Pedidos_Mes>0)", "semi_aditiva", "todas"],
    ], columns=["metrica_oficial", "tabela_proprietaria", "definicao", "aditividade", "areas_consumidoras"])


def build_reconciliation():
    return pd.DataFrame([
        ["REC001", "fato_item_pedido", "linhas", "6495", "igualdade", "critico"],
        ["REC002", "fato_item_pedido", "Bronze_Source_Row_Number distintos", "6495", "igualdade", "critico"],
        ["REC003", "fato_pedido", "linhas e Pedido_ID distintos", "2789", "igualdade", "critico"],
        ["REC004", "dim_cliente", "Cliente_ID distintos", "330", "igualdade", "critico"],
        ["REC005", "dim_produto", "Id_Item distintos", "36", "igualdade", "critico"],
        ["REC006", "fato_entrega", "linhas e Pedido_ID distintos", "2789", "igualdade", "critico"],
        ["REC007", "fato_item_pedido x fato_pedido", "soma Valor_Compra", "mesmo valor", "tolerancia 0.01", "critico"],
        ["REC008", "fato_item_pedido x fato_pedido", "soma Quantidade", "mesmo valor", "igualdade", "critico"],
        ["REC009", "todas as fatos", "FKs sem correspondencia", "0", "igualdade", "critico"],
        ["REC010", "dim_data", "dias de 2025", "365", "igualdade", "critico"],
        ["REC011", "dim_horario", "minutos do dia", "1440", "igualdade", "critico"],
        ["REC012", "fato_pedido", "Cupom_Status Nao informado", "5 pedidos", "igualdade", "critico"],
        ["REC013", "fato_pedido", "Pontos_Fidelidade_Observados constantes dentro do pedido", "0 conflitos", "igualdade", "critico"],
        ["REC014", "fato_cliente_mes", "meses sem compra apos primeira compra", "preservados", "validacao temporal", "critico"],
        ["REC015", "dim_data", "datas oficiais 2025 carregadas", "feriados e pontos facultativos conforme fonte", "igualdade", "critico"],
        ["REC016", "dim_cliente", "conflitos de atributos por Cliente_ID", "0 conflitos antes da selecao SCD Tipo 1", "unicidade semantica por atributo", "critico"],
    ], columns=["regra_id", "objeto", "controle", "esperado", "comparacao", "severidade"])


def build_decisions():
    return pd.DataFrame([
        ["DM001", "granularidade", "medidas de item pertencem a fato_item_pedido", "evitar duplicacao por pedido"],
        ["DM002", "granularidade", "frete e prazo pertencem a fato_entrega", "evitar multiplicacao por itens"],
        ["DM003", "financeiro", "nao criar fato_financeira duplicada", "usar medidas economicas das fatos existentes"],
        ["DM004", "marketing", "nao criar fato de interacao sem impressoes/cliques/leads", "pedidos apenas carregam contexto de marketing"],
        ["DM005", "cliente", "atributos comportamentais estaveis recebidos usam sufixo Informado", "permitir comparacao com valores recalculados"],
        ["DM006", "nulos", "Cupom nulo vira Nao informado somente na Gold analitica", "nao confundir com Nao"],
        ["DM007", "nulos", "renda nula do estagiario permanece nula", "decisao Silver preservada"],
        ["DM008", "data", "feriados nacionais e pontos facultativos federais de 2025 usam fonte oficial e classificacoes separadas", "nao equiparar ponto facultativo a feriado"],
        ["DM009", "financeiro", "margem e nao lucro", "nao existem impostos e despesas completas"],
        ["DM010", "marketing", "associacao e nao causalidade/ROI", "nao existem custos nem exposicoes de campanha"],
        ["DM011", "sk", "SKs devem ser deterministicas e reproduziveis", "idempotencia entre execucoes"],
        ["DM012", "historico", "Gold atual sobrescreve de forma controlada e snapshots preservam versoes", "sem append duplicador"],
        ["DM013", "cliente", "dim_cliente inicia como SCD Tipo 1", "nao existe historico confiavel para fabricar versoes retroativas"],
        ["DM014", "fidelidade", "Pontos_Fidelidade pertence ao pedido como valor observado", "o valor varia entre pedidos do mesmo cliente"],
        ["DM015", "relacionamento", "fato_entrega nao possui FK para fato_pedido", "fatos compartilham Pedido_ID e dimensoes conformadas"],
        ["DM016", "cliente_mes", "gerar meses sem compra desde a primeira compra observada ate 2025-12-31", "permitir recencia, risco, inatividade e reativacao"],
        ["DM017", "cliente_mes", "Data_Corte_SK representa o ultimo dia do mes", "coerencia com metricas de fechamento"],
        ["DM018", "horario", "preservar Data_Hora_Compra completa na fato_pedido", "dim_horario no minuto nao deve eliminar os segundos originais"],
        ["DM019", "data", "preservar o periodo parcial dos pontos facultativos na dim_data", "evitar interpretar ate_14h ou apos_13h como dia inteiro"],
    ], columns=["decisao_id", "tema", "decisao", "justificativa"])


def build_calendar_2025():
    """Calendario federal oficial; nao inclui feriados estaduais/municipais."""
    source_9783 = "Portaria MGI nº 9.783, de 27/12/2024"
    source_3197 = "Portaria MGI nº 3.197, de 28/04/2025"
    rows = [
        ["2025-01-01", "Confraternização Universal", "feriado_nacional", "dia_inteiro", "nacional", source_9783],
        ["2025-03-03", "Carnaval", "ponto_facultativo_federal", "dia_inteiro", "administracao_publica_federal", source_9783],
        ["2025-03-04", "Carnaval", "ponto_facultativo_federal", "dia_inteiro", "administracao_publica_federal", source_9783],
        ["2025-03-05", "Quarta-feira de Cinzas", "ponto_facultativo_federal", "ate_14h", "administracao_publica_federal", source_9783],
        ["2025-04-17", "Quinta-feira Santa", "ponto_facultativo_federal", "dia_inteiro", "administracao_publica_federal", source_9783],
        ["2025-04-18", "Paixão de Cristo", "feriado_nacional", "dia_inteiro", "nacional", source_9783],
        ["2025-04-21", "Tiradentes", "feriado_nacional", "dia_inteiro", "nacional", source_9783],
        ["2025-05-01", "Dia Mundial do Trabalho", "feriado_nacional", "dia_inteiro", "nacional", source_9783],
        ["2025-06-19", "Corpus Christi", "ponto_facultativo_federal", "dia_inteiro", "administracao_publica_federal", source_9783],
        ["2025-06-20", "Ponto facultativo federal", "ponto_facultativo_federal", "dia_inteiro", "administracao_publica_federal", source_3197],
        ["2025-09-07", "Independência do Brasil", "feriado_nacional", "dia_inteiro", "nacional", source_9783],
        ["2025-10-12", "Nossa Senhora Aparecida", "feriado_nacional", "dia_inteiro", "nacional", source_9783],
        ["2025-10-28", "Dia do Servidor Público federal", "ponto_facultativo_federal", "dia_inteiro", "administracao_publica_federal", source_9783],
        ["2025-11-02", "Finados", "feriado_nacional", "dia_inteiro", "nacional", source_9783],
        ["2025-11-15", "Proclamação da República", "feriado_nacional", "dia_inteiro", "nacional", source_9783],
        ["2025-11-20", "Dia Nacional de Zumbi e da Consciência Negra", "feriado_nacional", "dia_inteiro", "nacional", source_9783],
        ["2025-12-24", "Véspera do Natal", "ponto_facultativo_federal", "apos_13h", "administracao_publica_federal", source_9783],
        ["2025-12-25", "Natal", "feriado_nacional", "dia_inteiro", "nacional", source_9783],
        ["2025-12-31", "Véspera do Ano Novo", "ponto_facultativo_federal", "apos_13h", "administracao_publica_federal", source_9783],
    ]
    result = pd.DataFrame(rows, columns=[
        "data", "nome_dia_especial", "tipo_dia_especial", "periodo",
        "escopo", "fonte_normativa",
    ])
    result["data"] = pd.to_datetime(result["data"])
    return result


def build_external_sources():
    return pd.DataFrame([
        [
            "CAL2025_01", "calendario_2025", "Portaria MGI nº 9.783, de 27/12/2024",
            "fonte_oficial_primaria", "feriados nacionais e pontos facultativos federais de 2025",
            "https://www.gov.br/gestao/pt-br/acesso-a-informacao/institucional/atos-normativos/2024/copy_of_2023-portarias",
            "validada", "nao inclui feriados estaduais ou municipais",
        ],
        [
            "CAL2025_02", "calendario_2025", "Portaria MGI nº 3.197, de 28/04/2025",
            "fonte_oficial_primaria", "alteracao do calendario federal de 2025",
            "https://www.gov.br/gestao/pt-br/acesso-a-informacao/institucional/atos-normativos/2025/2025-portarias",
            "validada", "complementa a Portaria MGI nº 9.783/2024",
        ],
    ], columns=[
        "fonte_id", "objeto_destino", "documento", "tipo_fonte", "cobertura",
        "url", "status_validacao", "limitacao",
    ])


def validate_silver(silver):
    missing = sorted(set(BUSINESS_COLUMNS) - set(silver.columns))
    checks = [
        ["arquivo_silver_existe", True, True],
        ["colunas_obrigatorias_ausentes", len(missing), 0],
        ["linhas_silver", len(silver), EXPECTED_ROWS],
        ["pedidos_distintos", silver["Pedido_ID"].nunique(), EXPECTED_ORDERS],
        ["clientes_distintos", silver["Cliente_ID"].nunique(), EXPECTED_CLIENTS],
        ["produtos_distintos", silver["Id_Item"].nunique(), EXPECTED_PRODUCTS],
    ]
    result = pd.DataFrame(checks, columns=["controle", "observado", "esperado"])
    result["status"] = result.apply(lambda row: "aprovado" if row["observado"] == row["esperado"] else "reprovado", axis=1)
    result["detalhes"] = ""
    if missing:
        result.loc[result["controle"].eq("colunas_obrigatorias_ausentes"), "detalhes"] = " | ".join(missing)
    return result


def validate_specification(tables, columns, relationships, metrics):
    duplicate_columns = columns.duplicated(["tabela", "coluna_gold"]).sum()
    unknown_tables = set(columns["tabela"]) - set(tables["tabela"])
    facts_without_metrics = set(tables.loc[tables["tipo_tabela"].str.startswith("fato"), "tabela"]) - set(metrics["tabela_proprietaria"])
    checks = pd.DataFrame([
        ["tabelas_definidas", len(tables), 14],
        ["dimensoes_definidas", tables["tipo_tabela"].eq("dimensao").sum(), 10],
        ["fatos_definidas", tables["tipo_tabela"].str.startswith("fato").sum(), 4],
        ["colunas_duplicadas_na_mesma_tabela", int(duplicate_columns), 0],
        ["tabelas_de_colunas_nao_cadastradas", len(unknown_tables), 0],
        ["fatos_sem_metrica_oficial", len(facts_without_metrics), 0],
        ["relacionamentos_sem_referencia", relationships["referencia"].isna().sum(), 0],
    ], columns=["controle", "observado", "esperado"])
    checks["status"] = checks.apply(lambda row: "aprovado" if row["observado"] == row["esperado"] else "reprovado", axis=1)
    return checks


def validate_calendar(calendar, sources):
    return pd.DataFrame([
        ["datas_especiais_cadastradas", len(calendar), 19],
        ["feriados_nacionais", calendar["tipo_dia_especial"].eq("feriado_nacional").sum(), 10],
        ["pontos_facultativos_federais", calendar["tipo_dia_especial"].eq("ponto_facultativo_federal").sum(), 9],
        ["datas_duplicadas", calendar["data"].duplicated().sum(), 0],
        ["fontes_oficiais_validadas", sources["status_validacao"].eq("validada").sum(), 2],
    ], columns=["controle", "observado", "esperado"]).assign(
        status=lambda frame: frame.apply(
            lambda row: "aprovado" if row["observado"] == row["esperado"] else "reprovado",
            axis=1,
        )
    )


def style_workbook(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_cells in sheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells[:200]]
            width = min(max(max(map(len, values), default=0) + 2, 12), 55)
            sheet.column_dimensions[column_cells[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def save_report(sheets):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    temporary = OUTPUT_DIR / f".{OUTPUT_PATH.stem}.{uuid.uuid4().hex}.tmp.xlsx"
    try:
        with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
            for name, frame in sheets:
                frame.to_excel(writer, sheet_name=name, index=False)
        style_workbook(temporary)
        try:
            temporary.replace(OUTPUT_PATH)
        except PermissionError as error:
            raise PermissionError(
                f"Feche '{OUTPUT_PATH.name}' no Excel e aguarde a sincronizacao do OneDrive."
            ) from error
    finally:
        if temporary.exists():
            temporary.unlink()


def main():
    print("Iniciando especificacao do modelo dimensional Gold...")
    if not SILVER_PATH.is_file():
        raise FileNotFoundError(f"Silver nao encontrada: {SILVER_PATH}")
    silver = pd.read_parquet(SILVER_PATH)
    silver_checks = validate_silver(silver)
    tables = build_tables()
    columns = build_columns()
    bus_matrix = build_bus_matrix()
    relationships = build_relationships()
    metrics = build_metrics()
    reconciliation = build_reconciliation()
    decisions = build_decisions()
    calendar = build_calendar_2025()
    external_sources = build_external_sources()
    specification_checks = validate_specification(tables, columns, relationships, metrics)
    calendar_checks = validate_calendar(calendar, external_sources)
    failed = (
        silver_checks["status"].eq("reprovado").any()
        or specification_checks["status"].eq("reprovado").any()
        or calendar_checks["status"].eq("reprovado").any()
    )
    summary = pd.DataFrame([{
        "executado_em_utc": datetime.now(timezone.utc).isoformat(),
        "etapa": "11_especificacao_modelo_dimensional",
        "sha256_silver": sha256_file(SILVER_PATH),
        "linhas_silver": len(silver),
        "dimensoes": int(tables["tipo_tabela"].eq("dimensao").sum()),
        "fatos": int(tables["tipo_tabela"].str.startswith("fato").sum()),
        "colunas_especificadas": len(columns),
        "relacionamentos": len(relationships),
        "metricas_oficiais": len(metrics),
        "datas_especiais_2025": len(calendar),
        "fontes_externas_validadas": len(external_sources),
        "resultado": "REPROVADO" if failed else "APROVADO",
        "dados_modificados": False,
    }])
    save_report([
        ("resumo_modelo", summary),
        ("controles_silver", silver_checks),
        ("controles_especificacao", specification_checks),
        ("controles_calendario", calendar_checks),
        ("tabelas_modelo", tables),
        ("matriz_barramento", bus_matrix),
        ("colunas_modelo", columns),
        ("relacionamentos", relationships),
        ("metricas_oficiais", metrics),
        ("regras_reconciliacao", reconciliation),
        ("decisoes_modelagem", decisions),
        ("calendario_2025", calendar),
        ("fontes_externas", external_sources),
    ])
    MANIFEST_DIR.mkdir(parents=True, exist_ok=True)
    manifest = {
        "executado_em_utc": summary.iloc[0]["executado_em_utc"],
        "etapa": "11_especificacao_modelo_dimensional",
        "status": "failed" if failed else "success",
        "silver": {"path": str(SILVER_PATH), "sha256": sha256_file(SILVER_PATH), "rows": len(silver)},
        "report": {"path": str(OUTPUT_PATH), "sha256": sha256_file(OUTPUT_PATH)},
        "model": {"dimensions": 10, "facts": 4, "tables": tables["tabela"].tolist()},
        "external_sources": external_sources.to_dict(orient="records"),
        "calendar_2025": {
            "special_dates": len(calendar),
            "national_holidays": int(calendar["tipo_dia_especial"].eq("feriado_nacional").sum()),
            "federal_optional_days": int(calendar["tipo_dia_especial"].eq("ponto_facultativo_federal").sum()),
            "limitation": "nao inclui feriados estaduais ou municipais",
        },
    }
    MANIFEST_PATH.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"[OK] Relatorio gerado: {OUTPUT_PATH}")
    print(f"[OK] Manifesto gerado: {MANIFEST_PATH}")
    print(f"RESULTADO_FINAL: {'REPROVADO' if failed else 'APROVADO'}")
    if failed:
        raise ValueError("A especificacao ou a Silver nao passou nos controles obrigatorios.")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print("RESULTADO_FINAL: REPROVADO")
        print(f"ERRO: {type(error).__name__}: {error}")
        sys.exit(1)
